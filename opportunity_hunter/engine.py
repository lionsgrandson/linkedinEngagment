from __future__ import annotations

import base64
import hashlib
import html
import io
import ipaddress
import json
import os
import re
import smtplib
import socket
import ssl
import threading
import uuid
from datetime import datetime, timezone
from email.message import EmailMessage
from pathlib import Path
from typing import Any
from urllib.parse import quote, urlencode, urljoin, urlparse

import requests
from dotenv import load_dotenv


ROOT = Path(__file__).resolve().parents[1]
STATE_FILE = ROOT / "hunter_state.json"
load_dotenv(ROOT / ".env")

_STATE_LOCK = threading.RLock()

DEFAULT_STATE: dict[str, Any] = {
    "resume": {"name": "", "text": "", "updatedAt": ""},
    "jobs": [],
    "clients": [],
    "approvals": [],
}

USER_AGENT = "CodeCrafter-Opportunity-Hunter/1.0"
EMAIL_RE = re.compile(r"(?i)(?<![\w.+-])([a-z0-9._%+-]+@[a-z0-9.-]+\.[a-z]{2,})(?![\w.-])")
PHONE_RE = re.compile(r"(?<!\d)(\+?\d[\d\s().-]{6,}\d)(?!\d)")
TAG_RE = re.compile(r"<[^>]+>")
WS_RE = re.compile(r"\s+")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _clone(value: Any) -> Any:
    return json.loads(json.dumps(value, ensure_ascii=False))


def load_state() -> dict[str, Any]:
    with _STATE_LOCK:
        if not STATE_FILE.exists():
            return _clone(DEFAULT_STATE)
        try:
            saved = json.loads(STATE_FILE.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return _clone(DEFAULT_STATE)
        result = _clone(DEFAULT_STATE)
        for key in result:
            if key in saved:
                result[key] = saved[key]
        if not isinstance(result.get("jobs"), list):
            result["jobs"] = []
        if not isinstance(result.get("clients"), list):
            result["clients"] = []
        if not isinstance(result.get("approvals"), list):
            result["approvals"] = []
        return result


def save_state(state: dict[str, Any]) -> None:
    with _STATE_LOCK:
        temp = STATE_FILE.with_suffix(".tmp")
        temp.write_text(json.dumps(state, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        temp.replace(STATE_FILE)


def integration_status() -> dict[str, Any]:
    return {
        "ollama": {
            "configured": True,
            "url": os.getenv("OLLAMA_URL", "http://127.0.0.1:11434"),
            "model": os.getenv("OLLAMA_MODEL", "llama3.1:8b"),
        },
        "braveSearch": {"configured": bool(os.getenv("BRAVE_SEARCH_API_KEY"))},
        "adzuna": {
            "configured": bool(os.getenv("ADZUNA_APP_ID") and os.getenv("ADZUNA_APP_KEY")),
            "country": os.getenv("ADZUNA_COUNTRY", "gb"),
        },
        "smtp": {
            "configured": bool(os.getenv("OUTREACH_SMTP_HOST") and os.getenv("OUTREACH_SMTP_FROM")),
            "from": os.getenv("OUTREACH_SMTP_FROM", ""),
        },
        "twilio": {
            "configured": bool(
                os.getenv("TWILIO_ACCOUNT_SID")
                and os.getenv("TWILIO_AUTH_TOKEN")
                and os.getenv("TWILIO_FROM_NUMBER")
            ),
            "from": os.getenv("TWILIO_FROM_NUMBER", ""),
        },
    }


def _ollama_url() -> str:
    return os.getenv("OLLAMA_URL", "http://127.0.0.1:11434").rstrip("/")


def _ollama_model() -> str:
    return os.getenv("OLLAMA_MODEL", "llama3.1:8b")


def ollama_json(prompt: str, *, temperature: float = 0.15, timeout: int = 180) -> dict[str, Any]:
    payload = {
        "model": _ollama_model(),
        "prompt": prompt,
        "stream": False,
        "think": False,
        "format": "json",
        "options": {"temperature": temperature},
    }
    response = requests.post(f"{_ollama_url()}/api/generate", json=payload, timeout=timeout)
    response.raise_for_status()
    raw = str(response.json().get("response", "")).strip()
    if not raw:
        raise RuntimeError(f"Ollama model {_ollama_model()} returned an empty response")
    try:
        value = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise RuntimeError("Ollama did not return valid JSON") from exc
    if not isinstance(value, dict):
        raise RuntimeError("Ollama returned JSON, but not an object")
    return value


def ollama_text(prompt: str, *, temperature: float = 0.35, timeout: int = 180) -> str:
    payload = {
        "model": _ollama_model(),
        "prompt": prompt,
        "stream": False,
        "think": False,
        "options": {"temperature": temperature},
    }
    response = requests.post(f"{_ollama_url()}/api/generate", json=payload, timeout=timeout)
    response.raise_for_status()
    text = str(response.json().get("response", "")).strip()
    if not text:
        raise RuntimeError(f"Ollama model {_ollama_model()} returned an empty response")
    return text


def ollama_health() -> dict[str, Any]:
    try:
        response = requests.get(f"{_ollama_url()}/api/tags", timeout=5)
        response.raise_for_status()
        models = [str(item.get("name", "")) for item in response.json().get("models", []) if item.get("name")]
        return {"ok": True, "models": models, "selected": _ollama_model()}
    except (requests.RequestException, ValueError) as exc:
        return {"ok": False, "error": str(exc), "models": [], "selected": _ollama_model()}


def _extract_pdf(payload: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF resume support needs pypdf. Run: pip install -r requirements.txt") from exc
    reader = PdfReader(io.BytesIO(payload))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(payload: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX resume support needs python-docx. Run: pip install -r requirements.txt") from exc
    document = Document(io.BytesIO(payload))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_resume(filename: str, encoded: str) -> dict[str, Any]:
    if not filename:
        raise ValueError("Resume filename is required")
    raw = base64.b64decode(encoded, validate=True)
    if len(raw) > 8_000_000:
        raise ValueError("Resume file is too large; maximum size is 8 MB")
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".json", ".csv"}:
        text = raw.decode("utf-8", errors="replace")
    elif suffix == ".pdf":
        text = _extract_pdf(raw)
    elif suffix == ".docx":
        text = _extract_docx(raw)
    else:
        raise ValueError("Supported resume formats are PDF, DOCX, TXT, MD, JSON and CSV")
    text = text.replace("\x00", "").strip()
    if len(text) < 40:
        raise ValueError("The resume did not contain enough readable text")
    if len(text) > 60_000:
        text = text[:60_000]
    state = load_state()
    state["resume"] = {"name": filename, "text": text, "updatedAt": _now()}
    save_state(state)
    return {"name": filename, "characters": len(text), "preview": text[:800]}


def set_resume_text(name: str, text: str) -> dict[str, Any]:
    clean = str(text or "").replace("\x00", "").strip()
    if len(clean) < 40:
        raise ValueError("Paste at least a short resume/profile before saving")
    state = load_state()
    state["resume"] = {"name": name or "Pasted resume", "text": clean[:60_000], "updatedAt": _now()}
    save_state(state)
    return {"name": state["resume"]["name"], "characters": len(state["resume"]["text"])}


def resume_summary() -> dict[str, Any]:
    resume = load_state().get("resume", {})
    return {
        "name": str(resume.get("name", "")),
        "characters": len(str(resume.get("text", ""))),
        "updatedAt": str(resume.get("updatedAt", "")),
        "preview": str(resume.get("text", ""))[:600],
    }


def _safe_country(value: str) -> str:
    clean = re.sub(r"[^A-Za-z]", "", str(value or "")).lower()
    return clean[:2] if len(clean) >= 2 else ""


def brave_search(query: str, *, count: int = 10, country: str = "", language: str = "en") -> list[dict[str, Any]]:
    token = os.getenv("BRAVE_SEARCH_API_KEY", "").strip()
    if not token:
        raise RuntimeError("BRAVE_SEARCH_API_KEY is not configured")
    params: dict[str, Any] = {
        "q": query[:600],
        "count": max(1, min(20, int(count))),
        "search_lang": (language or "en")[:5],
        "extra_snippets": "true",
    }
    country_code = _safe_country(country)
    if country_code:
        params["country"] = country_code.upper()
    response = requests.get(
        "https://api.search.brave.com/res/v1/web/search",
        params=params,
        headers={"Accept": "application/json", "X-Subscription-Token": token, "User-Agent": USER_AGENT},
        timeout=25,
    )
    response.raise_for_status()
    rows = response.json().get("web", {}).get("results", [])
    results = []
    for row in rows:
        url = str(row.get("url", "")).strip()
        if not url.startswith(("http://", "https://")):
            continue
        snippets = [str(row.get("description", ""))]
        snippets.extend(str(item) for item in row.get("extra_snippets", []) if item)
        results.append({
            "title": str(row.get("title", "")).strip(),
            "url": url,
            "description": " ".join(item.strip() for item in snippets if item.strip())[:5000],
            "source": "brave",
        })
    return results


def adzuna_search_jobs(query: str, location: str, *, limit: int = 20, country: str = "") -> list[dict[str, Any]]:
    app_id = os.getenv("ADZUNA_APP_ID", "").strip()
    app_key = os.getenv("ADZUNA_APP_KEY", "").strip()
    if not app_id or not app_key:
        raise RuntimeError("ADZUNA_APP_ID and ADZUNA_APP_KEY are not configured")
    country_code = (_safe_country(country) or _safe_country(os.getenv("ADZUNA_COUNTRY", "gb")) or "gb")
    response = requests.get(
        f"https://api.adzuna.com/v1/api/jobs/{country_code}/search/1",
        params={
            "app_id": app_id,
            "app_key": app_key,
            "results_per_page": max(1, min(50, int(limit))),
            "what": query,
            "where": location,
            "content-type": "application/json",
        },
        headers={"Accept": "application/json", "User-Agent": USER_AGENT},
        timeout=25,
    )
    response.raise_for_status()
    jobs = []
    for row in response.json().get("results", []):
        jobs.append({
            "id": str(row.get("id", "")) or uuid.uuid4().hex,
            "title": str(row.get("title", "")).strip(),
            "company": str((row.get("company") or {}).get("display_name", "")).strip(),
            "location": str((row.get("location") or {}).get("display_name", "")).strip(),
            "description": str(row.get("description", "")).strip()[:8000],
            "url": str(row.get("redirect_url", "")).strip(),
            "created": str(row.get("created", "")),
            "salaryMin": row.get("salary_min"),
            "salaryMax": row.get("salary_max"),
            "source": "adzuna",
        })
    return jobs


def _web_results_to_jobs(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    jobs = []
    for row in rows:
        title = row.get("title", "")
        description = row.get("description", "")
        host = urlparse(row.get("url", "")).hostname or ""
        company = host.replace("www.", "")
        jobs.append({
            "id": hashlib.sha256(str(row.get("url", "")).encode()).hexdigest()[:18],
            "title": title,
            "company": company,
            "location": "",
            "description": description,
            "url": row.get("url", ""),
            "created": "",
            "salaryMin": None,
            "salaryMax": None,
            "source": "web",
        })
    return jobs


def _keyword_job_score(resume: str, job: dict[str, Any], target: str) -> int:
    resume_words = {word for word in re.findall(r"[a-zA-Z0-9+#.]{3,}", resume.lower())}
    job_words = set(re.findall(r"[a-zA-Z0-9+#.]{3,}", f"{job.get('title', '')} {job.get('description', '')}".lower()))
    target_words = set(re.findall(r"[a-zA-Z0-9+#.]{3,}", target.lower()))
    if not job_words:
        return 0
    overlap = len(resume_words & job_words) / max(1, min(40, len(job_words)))
    target_overlap = len(target_words & job_words) / max(1, len(target_words)) if target_words else 0
    return max(0, min(100, round(overlap * 80 + target_overlap * 20)))


def rank_jobs(jobs: list[dict[str, Any]], target: str) -> list[dict[str, Any]]:
    state = load_state()
    resume = str(state.get("resume", {}).get("text", ""))
    if not resume:
        for job in jobs:
            job["matchScore"] = _keyword_job_score(target, job, target)
            job["matchReason"] = "Resume not loaded; score is based on the requested job type only."
        return sorted(jobs, key=lambda item: item.get("matchScore", 0), reverse=True)

    compact = []
    for index, job in enumerate(jobs[:20]):
        compact.append({
            "index": index,
            "title": job.get("title", ""),
            "company": job.get("company", ""),
            "location": job.get("location", ""),
            "description": str(job.get("description", ""))[:1800],
        })
    prompt = f"""You are ranking job openings for one candidate. Be strict; do not inflate scores.
CANDIDATE RESUME:
{resume[:18000]}

TARGET ROLE:
{target}

JOBS:
{json.dumps(compact, ensure_ascii=False)}

Return JSON only with this exact shape:
{{"matches":[{{"index":0,"score":0,"reason":"one concise reason","gaps":["gap"]}}]}}
Score 0-100 based on actual evidence in the resume, not assumptions. Include every supplied index exactly once.
"""
    try:
        result = ollama_json(prompt, timeout=240)
        matches = {int(item.get("index")): item for item in result.get("matches", []) if isinstance(item, dict)}
    except Exception:
        matches = {}
    for index, job in enumerate(jobs):
        match = matches.get(index)
        if match:
            try:
                score = int(match.get("score", 0))
            except (TypeError, ValueError):
                score = 0
            job["matchScore"] = max(0, min(100, score))
            job["matchReason"] = str(match.get("reason", ""))[:500]
            job["gaps"] = [str(item)[:200] for item in match.get("gaps", [])[:8]]
        else:
            job["matchScore"] = _keyword_job_score(resume, job, target)
            job["matchReason"] = "Fallback keyword score; local AI ranking was unavailable for this result."
            job["gaps"] = []
    return sorted(jobs, key=lambda item: item.get("matchScore", 0), reverse=True)


def search_jobs(query: str, location: str, *, country: str = "", limit: int = 15, provider: str = "auto") -> dict[str, Any]:
    clean_query = str(query or "").strip()
    if not clean_query:
        raise ValueError("Enter the type of job you want")
    provider = str(provider or "auto").lower()
    errors: list[str] = []
    jobs: list[dict[str, Any]] = []
    used = ""
    if provider in {"auto", "adzuna"} and integration_status()["adzuna"]["configured"]:
        try:
            jobs = adzuna_search_jobs(clean_query, location, limit=limit, country=country)
            used = "adzuna"
        except Exception as exc:
            errors.append(f"Adzuna: {exc}")
            if provider == "adzuna":
                raise
    if not jobs and provider in {"auto", "brave"}:
        search = f'jobs "{clean_query}" {location} (site:jobs.lever.co OR site:boards.greenhouse.io OR site:jobs.workable.com OR site:careers-page.com OR site:linkedin.com/jobs)'
        rows = brave_search(search, count=min(20, limit), country=country)
        jobs = _web_results_to_jobs(rows)
        used = "brave"
    ranked = rank_jobs(jobs, clean_query)
    state = load_state()
    state["jobs"] = ranked[:100]
    save_state(state)
    return {"provider": used, "jobs": ranked, "errors": errors}


def make_application_pack(job: dict[str, Any]) -> dict[str, Any]:
    state = load_state()
    resume = str(state.get("resume", {}).get("text", ""))
    if not resume:
        raise ValueError("Upload or paste your resume first")
    prompt = f"""Create a truthful application pack for this job. Do not invent skills, dates, employers, degrees, metrics, or experience.
RESUME:
{resume[:20000]}

JOB:
{json.dumps(job, ensure_ascii=False)[:12000]}

Return JSON only:
{{"summary":"2-3 sentence fit summary","coverLetter":"concise tailored cover letter","likelyQuestions":[{{"question":"...","answer":"truthful answer based only on resume"}}],"keywords":["keyword"],"warnings":["anything the candidate should verify"]}}
"""
    result = ollama_json(prompt, timeout=240)
    approval = create_approval("job_application", {"job": job, "applicationPack": result})
    result["approvalId"] = approval["id"]
    return result


def _is_public_http_url(url: str) -> bool:
    try:
        parsed = urlparse(url)
        if parsed.scheme not in {"http", "https"} or not parsed.hostname:
            return False
        host = parsed.hostname.lower()
        if host in {"localhost", "localhost.localdomain"} or host.endswith(".local"):
            return False
        try:
            ip = ipaddress.ip_address(host)
            return not (ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved)
        except ValueError:
            return True
    except ValueError:
        return False


def _clean_html_text(raw: str) -> str:
    raw = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", raw)
    raw = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>", "\n", raw)
    raw = re.sub(r"(?i)</(?:p|div|li|h\d|section|article)>", "\n", raw)
    raw = TAG_RE.sub(" ", raw)
    raw = html.unescape(raw)
    lines = [WS_RE.sub(" ", line).strip() for line in raw.splitlines()]
    return "\n".join(line for line in lines if line)


def fetch_public_page(url: str, *, max_bytes: int = 2_000_000) -> dict[str, Any]:
    if not _is_public_http_url(url):
        raise ValueError("Only public HTTP(S) websites can be researched")
    response = requests.get(url, timeout=20, headers={"User-Agent": USER_AGENT}, allow_redirects=True)
    response.raise_for_status()
    if not _is_public_http_url(response.url):
        raise ValueError("Research request redirected to a non-public address")
    content_type = response.headers.get("Content-Type", "")
    if "html" not in content_type.lower() and "text" not in content_type.lower():
        raise ValueError("Research currently supports HTML/text pages only")
    raw = response.content[:max_bytes].decode(response.encoding or "utf-8", errors="replace")
    text = _clean_html_text(raw)
    emails = sorted(set(EMAIL_RE.findall(raw)))[:25]
    phones = []
    for match in PHONE_RE.findall(text):
        normalized = WS_RE.sub(" ", match).strip(" .,-()")
        digits = re.sub(r"\D", "", normalized)
        if 7 <= len(digits) <= 15:
            phones.append(normalized)
    phones = list(dict.fromkeys(phones))[:25]
    title_match = re.search(r"(?is)<title[^>]*>(.*?)</title>", raw)
    title = _clean_html_text(title_match.group(1))[:300] if title_match else ""
    links = re.findall(r"(?is)href=[\"']([^\"']+)[\"']", raw)
    contact_links = []
    for link in links:
        absolute = urljoin(response.url, html.unescape(link))
        lower = absolute.lower()
        if any(token in lower for token in ("contact", "about", "team", "company")) and _is_public_http_url(absolute):
            contact_links.append(absolute)
    return {
        "url": response.url,
        "title": title,
        "text": text[:30_000],
        "emails": emails,
        "phones": phones,
        "contactLinks": list(dict.fromkeys(contact_links))[:8],
    }


def search_clients(service_type: str, location: str, *, country: str = "", limit: int = 12) -> dict[str, Any]:
    service = str(service_type or "").strip()
    if not service:
        raise ValueError("Enter the type of client/business you want")
    query = f'"{service}" {location} company business contact -jobs -careers -linkedin -facebook -instagram'
    rows = brave_search(query, count=min(20, limit), country=country)
    clients = []
    for row in rows:
        host = urlparse(row["url"]).hostname or ""
        clients.append({
            "id": hashlib.sha256(row["url"].encode()).hexdigest()[:18],
            "name": row["title"],
            "website": row["url"],
            "domain": host.replace("www.", ""),
            "snippet": row["description"],
            "emails": [],
            "phones": [],
            "research": "",
            "fitScore": 0,
        })
    state = load_state()
    state["clients"] = clients
    save_state(state)
    return {"provider": "brave", "clients": clients}


def research_client(client: dict[str, Any], service_type: str) -> dict[str, Any]:
    page = fetch_public_page(str(client.get("website", "")))
    prompt = f"""Assess this company as a possible client for the offered service. Use only the supplied website evidence. Do not invent problems or claim you audited things that are not shown.
SERVICE OFFERED: {service_type}
COMPANY: {client.get('name', '')}
WEBSITE EXCERPT:
{page['text'][:14000]}

Return JSON only:
{{"fitScore":0,"reason":"why this is or is not a fit","research":"short factual company/relevance summary","angles":["specific outreach angle"],"risks":["uncertainty or reason not to contact"]}}
"""
    try:
        analysis = ollama_json(prompt, timeout=240)
    except Exception as exc:
        analysis = {"fitScore": 0, "reason": f"AI research unavailable: {exc}", "research": page["text"][:1000], "angles": [], "risks": []}
    enriched = {**client}
    enriched.update({
        "website": page["url"],
        "emails": page["emails"],
        "phones": page["phones"],
        "contactLinks": page["contactLinks"],
        "fitScore": max(0, min(100, int(analysis.get("fitScore", 0) or 0))),
        "fitReason": str(analysis.get("reason", ""))[:1000],
        "research": str(analysis.get("research", ""))[:5000],
        "angles": [str(item)[:500] for item in analysis.get("angles", [])[:10]],
        "risks": [str(item)[:500] for item in analysis.get("risks", [])[:10]],
    })
    state = load_state()
    state["clients"] = [enriched if item.get("id") == enriched.get("id") else item for item in state.get("clients", [])]
    save_state(state)
    return enriched


def _company_context(client: dict[str, Any]) -> str:
    return json.dumps({
        "name": client.get("name", ""),
        "website": client.get("website", ""),
        "snippet": client.get("snippet", ""),
        "research": client.get("research", ""),
        "angles": client.get("angles", []),
        "risks": client.get("risks", []),
    }, ensure_ascii=False)


def draft_client_email(client: dict[str, Any], service_type: str, *, sender_context: str = "") -> dict[str, Any]:
    prompt = f"""Draft a short first-contact business email. It must be honest, specific, low-pressure, and based only on the research below. Do not claim to have found technical problems unless the research explicitly proves them. Do not fake familiarity. Avoid spammy hype.
SERVICE: {service_type}
SENDER/COMPANY FACTS: {sender_context[:8000]}
CLIENT RESEARCH: {_company_context(client)[:12000]}
Return JSON only: {{"subject":"...","body":"...","reason":"why this angle fits"}}
"""
    draft = ollama_json(prompt)
    payload = {
        "client": client,
        "to": str((client.get("emails") or [""])[0]),
        "subject": str(draft.get("subject", ""))[:300],
        "body": str(draft.get("body", ""))[:12000],
        "reason": str(draft.get("reason", ""))[:1000],
    }
    approval = create_approval("email", payload)
    return {**payload, "approvalId": approval["id"]}


def draft_whatsapp(client: dict[str, Any], service_type: str, *, sender_context: str = "") -> dict[str, Any]:
    prompt = f"""Draft one concise first-contact WhatsApp message for a business prospect. Keep it professional and non-pushy. Use only verified research below. Do not invent familiarity, referrals, audits, urgency, discounts, or results.
SERVICE: {service_type}
SENDER/COMPANY FACTS: {sender_context[:6000]}
CLIENT RESEARCH: {_company_context(client)[:10000]}
Return JSON only: {{"message":"...","reason":"why this is relevant"}}
"""
    draft = ollama_json(prompt)
    payload = {
        "client": client,
        "phone": str((client.get("phones") or [""])[0]),
        "message": str(draft.get("message", ""))[:4000],
        "reason": str(draft.get("reason", ""))[:1000],
    }
    approval = create_approval("whatsapp", payload)
    return {**payload, "approvalId": approval["id"]}


def draft_call(client: dict[str, Any], service_type: str, *, sender_context: str = "") -> dict[str, Any]:
    prompt = f"""Write a short outbound business call script that will be spoken by text-to-speech. It must identify the caller/business clearly, state the reason for calling, be respectful, and avoid deceptive claims or fake urgency. Keep it under 110 words.
SERVICE: {service_type}
SENDER/COMPANY FACTS: {sender_context[:6000]}
CLIENT RESEARCH: {_company_context(client)[:10000]}
Return JSON only: {{"script":"...","reason":"why this script fits"}}
"""
    draft = ollama_json(prompt)
    payload = {
        "client": client,
        "phone": str((client.get("phones") or [""])[0]),
        "script": str(draft.get("script", ""))[:3000],
        "reason": str(draft.get("reason", ""))[:1000],
    }
    approval = create_approval("call", payload)
    return {**payload, "approvalId": approval["id"]}


def create_approval(kind: str, payload: dict[str, Any]) -> dict[str, Any]:
    state = load_state()
    approval = {
        "id": uuid.uuid4().hex,
        "kind": kind,
        "createdAt": _now(),
        "consumedAt": "",
        "payload": payload,
    }
    state["approvals"].append(approval)
    state["approvals"] = state["approvals"][-200:]
    save_state(state)
    return approval


def list_approvals(*, include_consumed: bool = False) -> list[dict[str, Any]]:
    rows = load_state().get("approvals", [])
    if include_consumed:
        return rows
    return [item for item in rows if not item.get("consumedAt")]


def _take_approval(approval_id: str, expected_kind: str) -> dict[str, Any]:
    state = load_state()
    for item in state.get("approvals", []):
        if item.get("id") != approval_id:
            continue
        if item.get("kind") != expected_kind:
            raise ValueError("Approval does not match this action")
        if item.get("consumedAt"):
            raise ValueError("Approval was already used")
        item["consumedAt"] = _now()
        save_state(state)
        return item
    raise ValueError("Approval was not found")


def _restore_approval(approval_id: str) -> None:
    state = load_state()
    for item in state.get("approvals", []):
        if item.get("id") == approval_id:
            item["consumedAt"] = ""
            save_state(state)
            return


def send_approved_email(approval_id: str, *, to_override: str = "") -> dict[str, Any]:
    approval = _take_approval(approval_id, "email")
    payload = approval["payload"]
    recipient = str(to_override or payload.get("to", "")).strip()
    if not EMAIL_RE.fullmatch(recipient):
        _restore_approval(approval_id)
        raise ValueError("Enter a valid recipient email address")
    host = os.getenv("OUTREACH_SMTP_HOST", "").strip()
    port = int(os.getenv("OUTREACH_SMTP_PORT", "587"))
    username = os.getenv("OUTREACH_SMTP_USER", "").strip()
    password = os.getenv("OUTREACH_SMTP_PASSWORD", "")
    sender = os.getenv("OUTREACH_SMTP_FROM", "").strip()
    if not host or not sender:
        _restore_approval(approval_id)
        raise RuntimeError("SMTP is not configured. Set OUTREACH_SMTP_HOST and OUTREACH_SMTP_FROM.")
    message = EmailMessage()
    message["From"] = sender
    message["To"] = recipient
    message["Subject"] = str(payload.get("subject", ""))[:300]
    message.set_content(str(payload.get("body", "")))
    try:
        if os.getenv("OUTREACH_SMTP_SSL", "0") == "1":
            with smtplib.SMTP_SSL(host, port, context=ssl.create_default_context(), timeout=20) as server:
                if username:
                    server.login(username, password)
                server.send_message(message)
        else:
            with smtplib.SMTP(host, port, timeout=20) as server:
                server.ehlo()
                if os.getenv("OUTREACH_SMTP_STARTTLS", "1") != "0":
                    server.starttls(context=ssl.create_default_context())
                    server.ehlo()
                if username:
                    server.login(username, password)
                server.send_message(message)
    except Exception:
        _restore_approval(approval_id)
        raise
    return {"sent": True, "to": recipient, "subject": message["Subject"]}


def approved_whatsapp_link(approval_id: str, *, phone_override: str = "") -> dict[str, Any]:
    approval = _take_approval(approval_id, "whatsapp")
    payload = approval["payload"]
    phone = str(phone_override or payload.get("phone", ""))
    digits = re.sub(r"\D", "", phone)
    if len(digits) < 7 or len(digits) > 15:
        _restore_approval(approval_id)
        raise ValueError("Enter the WhatsApp number including country code")
    message = str(payload.get("message", "")).strip()
    if not message:
        _restore_approval(approval_id)
        raise ValueError("Approved WhatsApp draft is empty")
    return {
        "opened": False,
        "url": f"https://wa.me/{digits}?text={quote(message)}",
        "note": "The approved message is prefilled. WhatsApp still requires the final send action in its UI.",
    }


def start_approved_call(approval_id: str, *, phone_override: str = "") -> dict[str, Any]:
    approval = _take_approval(approval_id, "call")
    payload = approval["payload"]
    sid = os.getenv("TWILIO_ACCOUNT_SID", "").strip()
    token = os.getenv("TWILIO_AUTH_TOKEN", "")
    from_number = os.getenv("TWILIO_FROM_NUMBER", "").strip()
    to_number = str(phone_override or payload.get("phone", "")).strip()
    if not sid or not token or not from_number:
        _restore_approval(approval_id)
        raise RuntimeError("Twilio is not configured. Set TWILIO_ACCOUNT_SID, TWILIO_AUTH_TOKEN and TWILIO_FROM_NUMBER.")
    if not to_number:
        _restore_approval(approval_id)
        raise ValueError("Enter a phone number for the approved call")
    script = str(payload.get("script", "")).strip()
    if not script:
        _restore_approval(approval_id)
        raise ValueError("Approved call script is empty")
    language = os.getenv("TWILIO_TTS_LANGUAGE", "en-US")
    voice = os.getenv("TWILIO_TTS_VOICE", "alice")
    twiml = f'<Response><Say language="{html.escape(language)}" voice="{html.escape(voice)}">{html.escape(script)}</Say></Response>'
    try:
        response = requests.post(
            f"https://api.twilio.com/2010-04-01/Accounts/{sid}/Calls.json",
            data={"To": to_number, "From": from_number, "Twiml": twiml},
            auth=(sid, token),
            timeout=25,
        )
        response.raise_for_status()
        result = response.json()
    except Exception:
        _restore_approval(approval_id)
        raise
    return {"started": True, "sid": result.get("sid"), "status": result.get("status"), "to": to_number}


def state_snapshot() -> dict[str, Any]:
    state = load_state()
    return {
        "resume": resume_summary(),
        "jobs": state.get("jobs", []),
        "clients": state.get("clients", []),
        "approvals": list_approvals(),
        "integrations": integration_status(),
        "ollama": ollama_health(),
    }
